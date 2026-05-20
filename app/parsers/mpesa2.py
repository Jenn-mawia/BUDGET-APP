"""M-PESA statement parser — Extract / Transform pipeline.

Reads an M-PESA .xlsx statement and returns normalized transactions.

Schema (shared by all parsers):
    date, description, amount (positive), direction ("in"/"out"),
    balance, source, account, is_overdraft_helper
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl


# --- Extract -----------------------------------------------------------

def extract(file_path: str | Path) -> list[tuple]:
    """Read the raw .xlsx rows. No cleaning here."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    return rows


# --- Transform ---------------------------------------------------------

def _to_amount(value) -> float:
    """Cell -> positive float. Handles floats, '1,234.56' strings, blanks."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return abs(float(value))
    text = str(value).strip().replace(",", "")
    try:
        return abs(float(text))
    except ValueError:
        return 0.0


def _header_index(rows: list[tuple]) -> int:
    """Index of the 'Receipt No.' row that starts the detailed table."""
    for i, row in enumerate(rows):
        if row[0] and str(row[0]).strip().lower() == "receipt no.":
            return i
    raise ValueError("Not a recognised M-PESA statement: no 'Receipt No.' header.")


def transform(rows: list[tuple], account: str) -> list[dict]:
    """Clean raw rows into normalized transaction dicts."""
    txns: list[dict] = []
    for row in rows[_header_index(rows) + 1:]:
        # A genuine transaction has a receipt no. and a real timestamp.
        if not row[0] or not isinstance(row[1], dt.datetime):
            continue

        # Details text is split across columns 3 and 5 by the converter.
        description = " ".join(
            str(row[c]).strip() for c in (3, 5) if row[c]
        )

        withdrawn, paid_in = _to_amount(row[10]), _to_amount(row[9])
        amount = withdrawn or paid_in
        if amount == 0.0:
            continue

        txns.append({
            "date": row[1].date(),
            "description": description,
            "amount": amount,
            "direction": "out" if withdrawn else "in",
            "balance": _to_amount(row[12]) if row[12] is not None else None,
            "source": "mpesa",
            "account": account,
            # Fuliza offset rows — real payment counted elsewhere.
            "is_overdraft_helper": description.lower().startswith("overdraft"),
        })
    return txns


# --- Orchestrator ------------------------------------------------------

def run(file_path: str | Path, account: str | None = None) -> list[dict]:
    """Extract + Transform an M-PESA statement into normalized transactions."""
    file_path = Path(file_path)
    rows = extract(file_path)
    return transform(rows, account or file_path.stem)
