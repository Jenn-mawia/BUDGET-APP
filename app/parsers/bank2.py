"""Bank statement parser — Extract / Transform pipeline.

Reads an I&M bank .xlsx statement (PDF-converted) into normalized
transactions, using the same schema as the M-PESA parser.

The converter produces a messy layout: column positions drift between
files, and balance + narrative are mashed into one cell. So we detect
fields by content, not fixed positions, and infer direction from how
the running balance moves between rows.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl


# --- Extract -----------------------------------------------------------

def extract(file_path: str | Path) -> list[tuple]:
    """Read the raw .xlsx rows. No cleaning here."""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    return rows


# --- Transform ---------------------------------------------------------

def _to_amount(value) -> float:
    """Cell -> float. Handles floats and '1,234.56' strings."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", ""))
    except ValueError:
        return 0.0


def _split_balance_narrative(text: str) -> tuple[float | None, str]:
    """Split a 'balance Cr/Dr narrative' cell into (balance, description).

    Example: '17,013.40 Cr   254729920461/MPESA Payment to ...'
             -> (17013.40, '254729920461/MPESA Payment to ...')
    """
    text = str(text).replace("\\n", " ").strip()
    m = re.match(r"([\d,]+\.\d{2})\s*(Cr|Dr)?\s*(.*)", text, re.DOTALL)
    if not m:
        return None, text
    balance = _to_amount(m.group(1))
    description = " ".join(m.group(3).split())
    return balance, description


def _find_amount(row: tuple, date_col: int, bal_col: int) -> float:
    """Find the lone transaction-amount number between date and balance.

    The amount sits in a drifting column somewhere after the dates and
    before the balance string. It is the only stray numeric cell there.
    """
    for j in range(date_col + 1, bal_col):
        v = row[j]
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return abs(float(v))
    return 0.0


def transform(rows: list[tuple], account: str) -> list[dict]:
    """Clean raw rows into normalized transactions.

    Direction is inferred from balance movement: a drop vs the previous
    balance means money out, a rise means money in.
    """
    txns: list[dict] = []
    prev_balance: float | None = None

    for row in rows:
        # A transaction row has a real date in column 1.
        if not isinstance(row[1], dt.datetime):
            continue

        # Balance + narrative live in the last non-empty cell of the row.
        last = next(
            (row[j] for j in range(len(row) - 1, -1, -1) if row[j] is not None),
            None,
        )
        if last is None:
            continue
        balance, description = _split_balance_narrative(last)
        bal_col = max(j for j in range(len(row)) if row[j] is not None)

        # The 'B/F' row only sets the opening balance — not a transaction.
        if description.upper().startswith("B/F"):
            prev_balance = balance
            continue

        amount = _find_amount(row, date_col=1, bal_col=bal_col)
        if amount == 0.0:
            continue

        # Infer direction from balance movement (the ground truth).
        if prev_balance is not None and balance is not None:
            direction = "out" if balance < prev_balance else "in"
        else:
            direction = "out"  # fallback if a balance is missing

        prev_balance = balance

        txns.append({
            "date": row[1].date(),
            "description": description,
            "amount": amount,
            "direction": direction,
            "balance": balance,
            "source": "bank",
            "account": account,
            "is_overdraft_helper": False,  # not applicable to bank rows
        })
    return txns


# --- Orchestrator ------------------------------------------------------

def run(file_path: str | Path, account: str | None = None) -> list[dict]:
    """Extract + Transform a bank statement into normalized transactions."""
    file_path = Path(file_path)
    rows = extract(file_path)
    return transform(rows, account or file_path.stem)
