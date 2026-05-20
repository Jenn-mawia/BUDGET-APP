"""
M-PESA statement parser.

Reads an M-PESA .xlsx statement (the kind exported from the Safaricom
self-service portal and converted to Excel) and returns a list of
normalized transaction dictionaries.

Normalized transaction schema (shared with the bank parser):
    date        : datetime.date   -- transaction date
    description : str             -- the narrative / details text
    amount      : float           -- always positive (the magnitude)
    direction   : str             -- "in" or "out"
    balance     : float | None    -- running balance after the transaction
    source      : str             -- always "mpesa" here
    account     : str             -- which account/file the row came from
    is_overdraft_helper : bool    -- True for Fuliza "OverDraft" offset rows
                                     (kept for traceability, excluded later)
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

import openpyxl


# The header labels we expect in the "DETAILED STATEMENT" section.
# We search for "Receipt No." to locate where the real table starts,
# because the file has a SUMMARY block above it with a variable height.
_RECEIPT_LABEL = "receipt no."


def _find_header_row(rows: list[tuple]) -> int:
    """
    Locate the index of the detailed-statement header row.

    The M-PESA file starts with a SUMMARY section, then a row that says
    'DETAILED STATEMENT', then the real header row containing
    'Receipt No.', 'Completion Time', 'Details', etc.

    We scan every row and return the index of the first one whose first
    cell reads 'Receipt No.' (case-insensitive). Raising a clear error
    here is deliberate: if the format ever changes, we want a loud,
    understandable failure rather than silently parsing garbage.
    """
    for idx, row in enumerate(rows):
        first_cell = row[0]
        if first_cell is not None and str(first_cell).strip().lower() == _RECEIPT_LABEL:
            return idx
    raise ValueError(
        "Could not find the 'Receipt No.' header row. "
        "This file may not be a recognised M-PESA statement."
    )


def _clean_amount(value) -> float:
    """
    Convert a spreadsheet cell into a float magnitude.

    M-PESA amounts can arrive as floats already (e.g. -200.0) or as
    strings with thousands separators (e.g. '1,508.27'). Withdrawals are
    stored as negatives in the file. We return the ABSOLUTE value here;
    direction ('in' vs 'out') is decided separately by the caller, so the
    sign convention of the source file can never corrupt our totals.

    Empty / missing cells become 0.0.
    """
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return abs(float(value))
    # string case: strip spaces and thousands separators
    text = str(value).strip().replace(",", "")
    if text == "":
        return 0.0
    try:
        return abs(float(text))
    except ValueError:
        return 0.0


def parse_mpesa(file_path: str | Path, account_label: str | None = None) -> list[dict]:
    """
    Parse an M-PESA .xlsx statement into normalized transaction dicts.

    Parameters
    ----------
    file_path : str | Path
        Path to the M-PESA .xlsx statement.
    account_label : str, optional
        A human-readable label for this account/file. If not given, the
        file's stem (name without extension) is used.

    Returns
    -------
    list[dict]
        One dict per transaction, following the normalized schema
        described in the module docstring.

    Notes
    -----
    Column layout of the DETAILED STATEMENT section (0-indexed):
        col 0  : Receipt No.
        col 1  : Completion Time (a datetime)
        col 3  : Details part 1   <-- description is SPLIT across
        col 5  : Details part 2   <-- columns 3 and 5; we join them
        col 8  : Transaction Status
        col 9  : Paid In   (positive when money came in)
        col 10 : Withdrawn (negative in the source file)
        col 12 : Balance
    """
    file_path = Path(file_path)
    account = account_label or file_path.stem

    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header_idx = _find_header_row(rows)

    transactions: list[dict] = []

    # Real transactions start on the row AFTER the header.
    for row in rows[header_idx + 1:]:
        receipt = row[0]

        # Skip blank rows: a genuine transaction always has a receipt no.
        if receipt is None or str(receipt).strip() == "":
            continue

        completion_time = row[1]
        # Skip rows where the timestamp is missing or not a real datetime
        # (footer/marketing rows occasionally slip through with text here).
        if not isinstance(completion_time, _dt.datetime):
            continue

        # --- Join the split description (columns 3 and 5) ---
        part1 = "" if row[3] is None else str(row[3]).strip()
        part2 = "" if row[5] is None else str(row[5]).strip()
        description = " ".join(p for p in (part1, part2) if p)

        # --- Amounts: Paid In (col 9) and Withdrawn (col 10) ---
        paid_in = _clean_amount(row[9])
        withdrawn = _clean_amount(row[10])

        # Decide direction and magnitude. One of the two is normally zero.
        if withdrawn > 0:
            amount = withdrawn
            direction = "out"
        else:
            amount = paid_in
            direction = "in"

        # Skip zero-value rows entirely (no money actually moved).
        if amount == 0.0:
            continue

        balance = row[12]
        balance = _clean_amount(balance) if balance is not None else None

        # --- Flag Fuliza / OverDraft offset rows ---
        # These "OverDraft of Credit Party" lines are the loan side of a
        # Fuliza-funded payment that is ALREADY recorded separately.
        # Counting them would double the spending total, so we mark them.
        is_overdraft_helper = description.lower().startswith("overdraft")

        transactions.append({
            "date": completion_time.date(),
            "description": description,
            "amount": amount,
            "direction": direction,
            "balance": balance,
            "source": "mpesa",
            "account": account,
            "is_overdraft_helper": is_overdraft_helper,
        })

    return transactions
