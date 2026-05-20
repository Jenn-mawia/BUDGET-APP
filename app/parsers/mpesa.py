"""
M-PESA statement parser.
 
Reads an M-PESA .xlsx statement (the kind exported from the Safaricom self-service portal and converted to Excel) and returns a list of
normalized transaction dictionaries.

Normalized transaction schema (shared with the bank parser):
    date            : datetime.date   -- transaction date
    description     : str             -- the narrative / details text
    amount          : float           -- always positive (the magnitude)
    direction       : str             -- "in" or "out"
    balance         : float | None    -- running balance after the transaction
    source          : str             -- always "mpesa" here
    account         : str             -- which account/file the row came from
    is_loan_advance : bool            -- True for Fuliza "OverDraft" offset rows (kept for traceability, excluded later)
"""

from __future__ import annotations
 
import datetime as _dt
from pathlib import Path
 
import openpyxl


## EXTRACT
def extract(file_path:str | Path)->list[tuple]:
    """
    Extracts rows from an M-PESA statement file.
    Args:
        file_path: path to the .xlsx file to extract from.
    Returns:
        A list of tuples, each tuple representing a row of data extracted from the file.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    return rows


## TRANSFORM
def _to_amount(value)->float:
    """
    Converts the amount value to a positive float representing an amount.Handles various formats including numeric types, strings with commas, and blanks.
    Args:
        value: The cell value to convert.
    Returns:
        A positive float representing the amount. Returns 0.0 for invalid or blank inputs.
    """
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return abs(float(value))
    text = str(value).strip().replace(",", "")
    try:
        return abs(float(text))
    except ValueError:
        return 0.0


def _header_index(rows:list[tuple])->int:
    """
    Finds the index of the header row in the extracted data, which is identified by the presence of "Receipt No." in the first column.
    Args:
        rows: A list of tuples representing the extracted rows from the M-PESA statement.
    Returns:
        The index of the header row. Raises a ValueError if no header is found.
    """
    for i, row in enumerate(rows):
        if row[0] and str(row[0]).strip().lower() == "receipt no.":
            return i
    raise ValueError("Not a recognised M-PESA statement: no 'Receipt No.' header.")

def transform(rows:list[tuple], account:str)->list[dict]:
    """
    Transforms the extracted rows into a list of normalized transaction dictionaries.
    Args:
        rows: A list of tuples representing the extracted rows from the M-PESA statement.
        account: The account identifier to associate with each transaction.
    Returns:
        A list of dictionaries, each representing a normalized transaction with fields such as date, description, amount, direction, balance, source, account, and is_loan_advance.
    """
    transactions:list[dict] = []
    for row in rows[_header_index(rows) + 1:]:
        if not row[0] or not isinstance(row[1], _dt.datetime): #  A genuine transaction has a receipt no. and a real timestamp.
            continue

        description = " ".join(str(row[c]).strip() for c in (3, 5) if row[c])      # Details text is split across columns 3 and 5 by the converter.
        
        withdrawn, paid_in = _to_amount(row[10]), _to_amount(row[9])
        
        amount = withdrawn or paid_in
        
        if amount == 0.0:
            continue

        transactions.append({
            "date": row[1].date(),
            "description": description,
            "amount": amount,
            "direction": "out" if withdrawn else "in",
            "balance": _to_amount(row[12]) if row[12] is not None else None,
            "source": "mpesa",
            "account": account,
            "is_loan_advance": description.lower().startswith("overdraft"), # Fuliza offset rows — real payment counted elsewhere.        
        })
    return transactions


## ORCHESTRATOR
def parse(file_path:str | Path, account:str | None=None)->list[dict]:
    """
    Orchestrates the parsing of an M-PESA statement file by extracting rows and transforming them into normalized transactions.
    Args:
        file_path: The path to the .xlsx file to parse.
        account: The account identifier to associate with each transaction.
    Returns:
        A list of dictionaries, each representing a normalized transaction extracted and transformed from the M-PESA statement.
    """
    file_path = Path(file_path)
    rows = extract(file_path)
    return transform(rows, account or file_path.stem)

